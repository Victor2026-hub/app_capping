"""Exportación a Excel, histograma SVG y recarga de datos (Entrega C)."""
import hashlib
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import Muestra
from app.services import consultas
from app.services.histograma import calcular_clases
from tests.conftest import INSUMOS

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _descargar(client, zona="Todas", metodo="percentil", parametro="95"):
    resp = client.get(f"/exportar?zona={zona}&metodo={metodo}&parametro={parametro}")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    ws = wb["Leyes_Capeadas"]
    enc = [c.value for c in ws[1]]
    filas = [dict(zip(enc, [c.value for c in f])) for f in ws.iter_rows(min_row=2)]
    resumen = {f[0].value: f[1].value for f in wb["Resumen"].iter_rows(min_row=2)}
    return resp, filas, resumen


def _verdaderos(filas):
    return [f for f in filas if f["Es_Outlier"] == "VERDADERO"]


# ---------------------------------------------------------------- Excel
def test_exportar_respuesta_y_filas(client):
    resp, filas, _ = _descargar(client)
    assert resp.mimetype == XLSX
    assert "Leyes_Capeadas.xlsx" in resp.headers["Content-Disposition"]
    assert "attachment" in resp.headers["Content-Disposition"]
    assert len(filas) == 174
    assert len(_verdaderos(filas)) == 9
    assert {f["Es_Outlier"] for f in filas} == {"VERDADERO", "FALSO"}
    assert sum(f["Es_Outlier"] == "FALSO" for f in filas) == 165


def test_exportar_hojas_y_encabezados(client):
    resp = client.get("/exportar")
    wb = load_workbook(BytesIO(resp.data))
    assert wb.sheetnames == ["Leyes_Capeadas", "Resumen"]
    enc = [c.value for c in wb["Leyes_Capeadas"][1]]
    for col in ("Muestra_ID", "Au_gpt", "Au_usado", "Es_Outlier", "Au_capped"):
        assert col in enc


def test_exportar_capped_maximo_y_sin_exceder_umbral(client):
    _, filas, resumen = _descargar(client)
    capped = [f["Au_capped"] for f in filas]
    assert max(capped) == pytest.approx(8.36, abs=0.01)
    umbral = resumen["Umbral (g/t)"]
    assert all(c <= umbral + 1e-9 for c in capped)


def test_exportar_conserva_ley_original(client):
    _, filas, _ = _descargar(client)
    por_id = {f["Muestra_ID"]: f for f in filas}
    assert por_id["M-4102"]["Au_gpt"] == -99
    assert por_id["M-4128"]["Au_gpt"] == -99
    assert por_id["M-4102"]["Au_usado"] == pytest.approx(0.01)
    assert por_id["M-4026"]["Au_gpt"] == pytest.approx(95.40)
    assert por_id["M-4026"]["Au_capped"] == pytest.approx(8.36, abs=0.01)


def test_exportar_outliers_son_las_9_de_mayor_ley(client):
    _, filas, _ = _descargar(client)
    top9 = sorted(filas, key=lambda f: f["Au_usado"], reverse=True)[:9]
    assert {f["Muestra_ID"] for f in top9} == {f["Muestra_ID"] for f in _verdaderos(filas)}


def test_exportar_hoja_resumen(client):
    _, _, r = _descargar(client)
    assert r["Zona"] == "Todas"
    assert r["Método"] == "Percentil"
    assert r["Parámetro"] == "P = 95"
    assert r["n (muestras)"] == 174
    assert r["Umbral (g/t)"] == pytest.approx(8.36, abs=0.01)
    assert r["Outliers"] == 9
    assert r["Media capeada (g/t)"] == pytest.approx(2.23, abs=0.01)
    assert r["Máximo: Muestra_ID"] == "M-4026"


@pytest.mark.parametrize("zona,n,outs,umbral", [
    ("Veta Norte", 122, 7, 8.83),
    ("Veta Sur", 52, 3, 6.02),
])
def test_exportar_por_zona(client, zona, n, outs, umbral):
    _, filas, r = _descargar(client, zona=zona.replace(" ", "%20"))
    assert len(filas) == n
    assert len(_verdaderos(filas)) == outs
    assert {f["Zona"] for f in filas} == {zona}
    assert r["Zona"] == zona
    assert r["Umbral (g/t)"] == pytest.approx(umbral, abs=0.01)


def test_exportar_media_k_sigma_2(client):
    _, filas, r = _descargar(client, metodo="media_k_sigma", parametro="2")
    assert len(filas) == 174
    assert len(_verdaderos(filas)) == 6
    assert r["Umbral (g/t)"] == pytest.approx(22.09, abs=0.01)
    assert r["Parámetro"] == "k = 2"
    assert r["Media capeada (g/t)"] == pytest.approx(2.76, abs=0.01)


def test_exportar_parametro_invalido_400(client):
    assert client.get("/exportar?zona=Todas&metodo=percentil&parametro=50").status_code == 400


# ------------------------------------------------------------ Histograma
def test_clases_cuentas_suman_len_de_cada_lista():
    orig = [0.01, 0.5, 1.0, 2.0, 7.0, 95.4, 3.3]
    cap = [0.01, 0.5, 1.0, 2.0, 7.0, 8.36]
    c = calcular_clases(orig, cap)
    assert sum(c["original"]) == len(orig)
    assert sum(c["capeada"]) == len(cap)
    assert len(c["bordes"]) == len(c["original"]) + 1 == len(c["capeada"]) + 1


def test_clases_misma_particion_para_ambas_series():
    c = calcular_clases([0.01, 100.0], [0.5, 5.0])
    assert len(c["original"]) == len(c["capeada"])
    assert c["bordes"] == sorted(c["bordes"])
    assert c["bordes"][0] <= 0.01 and c["bordes"][-1] >= 100.0


def test_clases_valor_igual_a_borde_cae_en_una_sola_clase():
    # 1, 10 y 100 son bordes exactos de la partición (décadas completas)
    c = calcular_clases([1.0, 10.0, 100.0], [1.0, 10.0, 100.0])
    assert sum(c["original"]) == 3
    assert max(c["original"]) == 1  # ninguno se cuenta doble ni en la misma clase
    assert sum(c["capeada"]) == 3


def test_clases_valor_en_borde_intermedio():
    c = calcular_clases([1.0, 100.0], [1.0, 100.0])
    borde = c["bordes"][1]  # primer borde interior (~1.585)
    c2 = calcular_clases([1.0, borde, 100.0], [1.0, 100.0])
    assert sum(c2["original"]) == 3
    assert c2["bordes"] == c["bordes"]


def test_histograma_app_real_cuentas_suman_174(app):
    with app.app_context():
        d = consultas.analisis("Todas", "percentil", 95)
    c = calcular_clases(d["originales"], d["capeadas"])
    assert sum(c["original"]) == 174
    assert sum(c["capeada"]) == 174


def test_home_svg_despues_de_tabla_outliers_y_seguro(client):
    html = client.get("/").get_data(as_text=True)
    pos_tabla = html.index("<h2>Outliers</h2>")
    pos_svg = html.index("<svg", pos_tabla)
    assert pos_svg > pos_tabla
    svg = html[pos_svg:html.index("</svg>", pos_svg) + 6]
    assert "<script" not in svg.lower()
    assert svg.replace('xmlns="http://www.w3.org/2000/svg"', "").count("http") == 0


# --------------------------------------------------------------- Recarga
def _sha_insumos():
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(INSUMOS.glob("*.xlsx"))}


def test_recargar_restaura_datos_y_no_toca_insumos(app, client):
    antes = _sha_insumos()
    assert len(antes) == 2
    with app.app_context():
        for m in db.session.scalars(db.select(Muestra).limit(20)).all():
            db.session.delete(m)
        db.session.commit()
        assert consultas.contadores()["muestras"] == 154

    resp = client.post("/recargar", data={"zona": "Todas", "metodo": "percentil", "parametro": "95"})
    assert resp.status_code == 302

    html = client.get(resp.headers["Location"]).get_data(as_text=True)
    assert "Datos recargados" in html
    with app.app_context():
        assert consultas.contadores() == {"sondajes": 12, "muestras": 174, "inconsistencias": 3}
    assert _sha_insumos() == antes


def test_recargar_get_405(client):
    assert client.get("/recargar").status_code == 405


def test_recargar_con_insumos_inexistente_da_302_con_flash_de_error(app, client, tmp_path):
    vacia = tmp_path / "vacia"
    vacia.mkdir()
    app.config["INSUMOS_DIR"] = str(vacia)
    resp = client.post("/recargar", data={})
    assert resp.status_code == 302  # no 500
    html = client.get(resp.headers["Location"]).get_data(as_text=True)
    assert "No se pudo recargar los datos" in html
    assert 'class="flash error"' in html
    # Estado de la BD tras la falla (solo se documenta: lee Excel antes de drop_all)
    with app.app_context():
        print("ESTADO BD TRAS FALLA:", consultas.contadores())
