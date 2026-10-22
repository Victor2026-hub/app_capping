"""Genera Leyes_Capeadas.xlsx en memoria."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.services import capping, consultas


def _encabezado(ws, columnas, anchos):
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"


def generar_excel(zona: str, metodo: str, parametro) -> BytesIO:
    muestras = consultas.muestras_por_zona(zona)  # ya ordenadas por Muestra_ID
    leyes = [m["au_gpt"] for m in muestras]
    ids = [(m["muestra_id"], m["sondaje_id"]) for m in muestras]
    res = capping.resumen(leyes, metodo, parametro, ids)
    aplicado = capping.aplicar(leyes, res["umbral"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Leyes_Capeadas"
    _encabezado(
        ws,
        ["Muestra_ID", "Sondaje_ID", "Desde_m", "Hasta_m", "Au_gpt", "Lote_Laboratorio",
         "Zona", "Au_usado", "Es_Outlier", "Au_capped"],
        [13, 13, 10, 10, 10, 18, 13, 11, 12, 11],
    )
    for m, r in zip(muestras, aplicado):
        ws.append([
            m["muestra_id"], m["sondaje_id"], m["desde_m"], m["hasta_m"], m["au_original"],
            m["lote"], m["zona"], r["au_gpt"], "VERDADERO" if r["es_outlier"] else "FALSO",
            r["au_capped"],
        ])
    for fila in ws.iter_rows(min_row=2):
        for idx in (2, 3, 4, 7, 9):
            fila[idx].number_format = "0.00"

    rs = wb.create_sheet("Resumen")
    _encabezado(rs, ["Indicador", "Valor"], [32, 34])
    etiqueta_param = ("P = " if metodo == capping.METODO_PERCENTIL else "k = ") + str(parametro)
    filas = [
        ("Zona", zona),
        ("Método", consultas.ETIQUETAS_METODO[metodo]),
        ("Parámetro", etiqueta_param),
        ("n (muestras)", res["n"]),
        ("Media (g/t)", res["media"]),
        ("Mediana (g/t)", res["mediana"]),
        ("Máximo (g/t)", res["maximo"]),
        ("Máximo: Muestra_ID", res["maximo_muestra_id"]),
        ("Máximo: Sondaje_ID", res["maximo_sondaje_id"]),
        ("Umbral (g/t)", res["umbral"]),
        ("Outliers", res["outliers"]),
        ("Media capeada (g/t)", res["media_capeada"]),
        ("Variación de la media (%)", res["variacion_pct"]),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for nombre, valor in filas:
        rs.append([nombre, valor])
        celda = rs.cell(row=rs.max_row, column=2)
        if isinstance(valor, float):
            celda.number_format = "0.00"
        celda.alignment = Alignment(horizontal="left")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
